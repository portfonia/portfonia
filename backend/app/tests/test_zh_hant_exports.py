"""Issue #585 acceptance tests; no LLM calls or production data."""

from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
import yaml
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models.holding import Holding
from app.models.user import User
from app.services import holding_parser, holdings_export, portfolio_export
from app.services.zh_hant import to_traditional
from app.tests.conftest import TEST_USER_ID, seed_user


@pytest.fixture()
def traditional_user(db_session: Session) -> User:
    user = seed_user(db_session, TEST_USER_ID)
    user.locale = "zh-Hant"
    db_session.flush()
    return user


def test_template_traditional(app_client: TestClient, traditional_user: User) -> None:
    response = app_client.get("/holdings/template?locale=zh-Hant")
    simplified = app_client.get("/holdings/template?locale=zh")
    assert response.status_code == simplified.status_code == 200
    assert response.text != simplified.text
    assert to_traditional(response.text) == response.text
    assert response.text.startswith("##### 持倉模板\n")
    assert app_client.get("/holdings/template").text == response.text
    assert (
        app_client.get("/holdings/template?locale=fr").text
        == app_client.get("/holdings/template?locale=en").text
    )


def test_export_traditional_user_and_override(
    app_client: TestClient, db_session: Session, traditional_user: User
) -> None:
    holding = Holding(
        user_id=TEST_USER_ID,
        name="Apple",
        ticker="AAPL",
        currency="USD",
        pricing_mode="auto",
        asset_type="stock",
        shares=Decimal("10"),
        avg_cost=Decimal("180"),
        broker="IBKR",
        account="IRA",
        portfolio="Growth",
        notes="core",
    )
    db_session.add(holding)
    db_session.flush()
    response = app_client.get("/holdings/export")
    assert response.status_code == 200
    assert response.text.startswith("##### 持倉模板\n")
    assert response.text == holdings_export.render_export([holding], "zh-Hant")
    for locale in ("en", "zh", "fr", "zh-Hant"):
        body = app_client.get(f"/holdings/export?locale={locale}").text
        assert body.split("\n\n")[-1] == response.text.split("\n\n")[-1]
        assert body.endswith(holdings_export.render_holding_line(holding) + "\n")
    assert app_client.get("/holdings/export?locale=en").text.startswith("##### Holdings template")
    assert (
        app_client.get("/holdings/export?locale=fr").text
        == app_client.get("/holdings/export?locale=en").text
    )


@pytest.mark.parametrize("fmt", ["md", "xlsx"])
@pytest.mark.parametrize("explicit", [True, False])
def test_portfolio_traditional_headers_labels_and_fallback(
    app_client: TestClient, traditional_user: User, fmt: str, explicit: bool
) -> None:
    query = "&locale=zh-Hant" if explicit else ""
    response = app_client.get(f"/portfolio/export?format={fmt}{query}")
    assert response.status_code == 200
    expected = [portfolio_export._HEADERS_ZH_HANT[c] for c in portfolio_export.EXPORT_COLUMNS]
    as_of = portfolio_export._AS_OF_LABEL_BY_LOCALE["zh-Hant"]
    currency = portfolio_export._BASE_CURRENCY_LABEL_BY_LOCALE["zh-Hant"]
    english = app_client.get(f"/portfolio/export?format={fmt}&locale=en")
    unknown = app_client.get(f"/portfolio/export?format={fmt}&locale=fr")
    if fmt == "md":
        lines = response.text.splitlines()
        assert [c.strip() for c in lines[3].strip("|").split("|")] == expected
        assert lines[0].startswith(as_of + ":")
        assert lines[1] == currency + ": USD"
        assert unknown.text == english.text
        assert lines[4:] == english.text.splitlines()[4:]
    else:

        def rows(content: bytes) -> list[tuple[object, ...]]:
            ws = load_workbook(BytesIO(content)).active
            assert ws is not None
            return list(ws.iter_rows(values_only=True))

        result = rows(response.content)
        assert list(result[3]) == expected
        assert result[0][0] == as_of
        assert result[1][:2] == (currency, "USD")
        assert rows(unknown.content) == rows(english.content)
        assert result[4:] == rows(english.content)[4:]


def test_parser_vocabulary_adds_ordered_unique_traditional_variants(tmp_path: Path) -> None:
    raw = yaml.safe_load(holding_parser._DEFAULT_VOCAB_FILE.read_text())
    raw["cny_institutions"] = ["建设证券", "中信", "建设证券"]
    raw["common_cn_platforms"] = ["支付宝", "微信"]
    path = tmp_path / "vocab.yml"
    path.write_text(yaml.safe_dump(raw, allow_unicode=True), encoding="utf-8")
    vocab = holding_parser._load_holding_parser_vocab(path)
    assert vocab.cny_institutions == ("建设证券", to_traditional("建设证券"), "中信")
    assert vocab.common_cn_platforms == ("支付宝", to_traditional("支付宝"), "微信")
    prompt = holding_parser._build_system_prompt(vocab)
    assert "建设证券" in prompt and to_traditional("建设证券") in prompt
    assert prompt.count("中信") == 1
    for key in (
        "futu",
        "stock_connect",
        "cash",
        "margin",
        "deposit",
        "money_market",
        "index_fund",
        "wmp_terms",
        "a_share_terms",
        "us_market_zh",
        "hk_market_zh",
    ):
        original = raw[key].split("/")
        expected = list(
            dict.fromkeys(term for word in original for term in (word, to_traditional(word)))
        )
        assert getattr(vocab, key).split("/") == expected
        for term in expected:
            assert term in prompt
    for alias, market in raw["market_aliases_zh"].items():
        assert vocab.market_aliases_zh[alias] == market
        assert vocab.market_aliases_zh[to_traditional(alias)] == market


def test_traditional_dialect_structure() -> None:
    assert portfolio_export._HEADERS_ZH_HANT.keys() == portfolio_export._HEADERS_EN.keys()
    assert len(holdings_export._RULES_ZH_HANT.splitlines()) == len(
        holdings_export._RULES_ZH.splitlines()
    )
    assert len(holdings_export._EXAMPLES_ZH_HANT.splitlines()) == len(
        holdings_export._EXAMPLES_ZH.splitlines()
    )
    for traditional, simplified in (
        (holdings_export._RULES_ZH_HANT, holdings_export._RULES_ZH),
        (holdings_export._EXAMPLES_ZH_HANT, holdings_export._EXAMPLES_ZH),
    ):
        for line, source in zip(traditional.splitlines(), simplified.splitlines(), strict=True):
            for field in (
                "account:",
                "portfolio:",
                "notes:",
                "asset_type:",
                "market:",
                "pricing_mode:",
            ):
                assert line.count(field) == source.count(field)
